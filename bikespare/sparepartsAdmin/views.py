from django.contrib import messages
from django.contrib.auth import login, logout
from django.contrib.auth.decorators import user_passes_test
from django.contrib.auth.forms import AuthenticationForm
from django.db.models import Q
from django.shortcuts import get_object_or_404, redirect, render

from inventory.models import Brand, Category, Enquiry, Part

from .forms import (
    BrandForm,
    CategoryForm,
    EnquiryForm,
    PartForm,
)


# ============================================================
# ADMIN AUTHENTICATION
# ============================================================

def is_admin(user):
    return user.is_authenticated and user.is_staff


admin_required = user_passes_test(
    is_admin,
    login_url="dashboard:admin_login"
)


def admin_login_view(request):

    # Already logged in as staff
    if request.user.is_authenticated and request.user.is_staff:
        return redirect("dashboard:admin_dashboard")

    if request.method == "POST":

        form = AuthenticationForm(
            request,
            data=request.POST
        )

        if form.is_valid():

            user = form.get_user()

            if user.is_staff:

                login(request, user)

                return redirect(
                    "dashboard:admin_dashboard"
                )

            form.add_error(
                None,
                "You do not have administrator permissions."
            )

    else:

        form = AuthenticationForm()

    return render(
        request,
        "dashboard/login.html",
        {
            "form": form
        }
    )


def admin_logout_view(request):

    logout(request)

    return redirect(
        "dashboard:admin_login"
    )


# ============================================================
# DASHBOARD
# ============================================================

@admin_required
def admin_dashboard(request):

    total_parts = Part.objects.count()

    low_stock_parts = Part.objects.filter(
        quantity__lte=5,
        quantity__gt=0
    ).count()

    out_of_stock_parts = Part.objects.filter(
        quantity=0
    ).count()

    total_enquiries = Enquiry.objects.count()

    pending_enquiries = Enquiry.objects.filter(
        is_handled=False
    ).count()

    query = request.GET.get(
        "q",
        ""
    ).strip()

    brand_filter = request.GET.get(
        "brand",
        ""
    ).strip()

    category_filter = request.GET.get(
        "category",
        ""
    ).strip()

    parts = (
        Part.objects
        .select_related(
            "brand",
            "category"
        )
        .all()
    )

    if query:

        parts = parts.filter(
            Q(name__icontains=query)
            | Q(part_number__icontains=query)
            | Q(compatible_with__icontains=query)
            | Q(description__icontains=query)
        )

    if brand_filter:

        parts = parts.filter(
            brand_id=brand_filter
        )

    if category_filter:

        parts = parts.filter(
            category_id=category_filter
        )

    enquiries = Enquiry.objects.all()[:10]

    brands = Brand.objects.all()

    categories = Category.objects.all()

    context = {

        "total_parts": total_parts,

        "low_stock_parts": low_stock_parts,

        "out_of_stock_parts": out_of_stock_parts,

        "total_enquiries": total_enquiries,

        "pending_enquiries": pending_enquiries,

        "parts": parts,

        "enquiries": enquiries,

        "brands": brands,

        "categories": categories,

        "search_query": query,

        "selected_brand": brand_filter,

        "selected_category": category_filter,
    }

    return render(
        request,
        "dashboard/admin_dashboard.html",
        context
    )


# ============================================================
# PART CRUD
# ============================================================
import os
from decimal import Decimal, InvalidOperation

from django.contrib import messages
from django.contrib.auth.decorators import user_passes_test
from django.db import transaction
from django.shortcuts import redirect, render
from django.utils.text import slugify

from openpyxl import load_workbook

from inventory.models import Part, Brand, Category

def clean_part_number(value):
    if value is None:
        return ""

    # Excel numeric value: 2.0 -> "2"
    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        return str(value).strip()

    # Integer
    if isinstance(value, int):
        return str(value)

    # Text
    value = str(value).strip()

    # Text value like "2.0" -> "2"
    try:
        decimal_value = Decimal(value)

        if decimal_value == decimal_value.to_integral_value():
            return str(int(decimal_value))

    except (InvalidOperation, ValueError):
        pass

    return value
import os
from decimal import Decimal, InvalidOperation

from django.contrib import messages
from django.db import transaction
from django.shortcuts import redirect
from django.utils.text import slugify

from openpyxl import load_workbook

from .models import Part, Brand, Category
# from .utils import clean_part_number  # adjust if your function is elsewhere


@admin_required
def part_import_view(request):

    if request.method != "POST":
        return redirect("dashboard:part_list")

    excel_file = request.FILES.get("file")

    if not excel_file:
        messages.error(request, "Please select an Excel file.")
        return redirect("dashboard:part_list")

    extension = os.path.splitext(excel_file.name)[1].lower()

    if extension not in [".xlsx", ".xlsm"]:
        messages.error(
            request,
            "Please upload an Excel .xlsx or .xlsm file."
        )
        return redirect("dashboard:part_list")

    workbook = None

    try:
        # IMPORTANT:
        # read_only=True prevents the entire Excel workbook
        # from being loaded into memory.
        workbook = load_workbook(
            excel_file,
            read_only=True,
            data_only=True,
        )

        worksheet = workbook.active

        # -----------------------------------------
        # Read header row
        # -----------------------------------------

        header_row = next(
            worksheet.iter_rows(
                min_row=1,
                max_row=1,
                values_only=True
            ),
            None
        )

        if not header_row:
            messages.error(
                request,
                "The Excel file is empty."
            )
            return redirect("dashboard:part_list")

        headers = [
            str(value).strip()
            if value is not None
            else ""
            for value in header_row
        ]

        required_columns = [
            "Part Number",
            "Item Name",
            "Brand",
            "Category",
            "Qty",
            "Discount (%)",
            "Amount",
        ]

        missing_columns = [
            column
            for column in required_columns
            if column not in headers
        ]

        if missing_columns:
            messages.error(
                request,
                "Missing columns: " + ", ".join(missing_columns)
            )
            return redirect("dashboard:part_list")

        column_index = {
            header: index
            for index, header in enumerate(headers)
        }

        created_count = 0
        updated_count = 0
        skipped_count = 0

        # Don't allow thousands of errors to accumulate in memory.
        errors = []

        # -----------------------------------------
        # Cache Brand and Category objects
        # -----------------------------------------

        brand_cache = {}
        category_cache = {}

        # -----------------------------------------
        # Process Excel row by row
        # -----------------------------------------

        for row_number, row in enumerate(
            worksheet.iter_rows(
                min_row=2,
                values_only=True
            ),
            start=2
        ):

            try:

                part_number = row[
                    column_index["Part Number"]
                ]

                name = row[
                    column_index["Item Name"]
                ]

                brand_name = row[
                    column_index["Brand"]
                ]

                category_name = row[
                    column_index["Category"]
                ]

                quantity = row[
                    column_index["Qty"]
                ]

                discount = row[
                    column_index["Discount (%)"]
                ]

                amount = row[
                    column_index["Amount"]
                ]

                # -----------------------------------------
                # Required fields
                # -----------------------------------------

                if not part_number:
                    raise ValueError(
                        "Part Number is empty"
                    )

                if not name:
                    raise ValueError(
                        "Item Name is empty"
                    )

                if not brand_name:
                    raise ValueError(
                        "Brand is empty"
                    )

                if not category_name:
                    raise ValueError(
                        "Category is empty"
                    )

                # -----------------------------------------
                # Clean strings
                # -----------------------------------------

                part_number = clean_part_number(
                    part_number
                )

                name = str(name).strip()

                brand_name = str(
                    brand_name
                ).strip()

                category_name = str(
                    category_name
                ).strip()

                # -----------------------------------------
                # Quantity
                # -----------------------------------------

                quantity = int(
                    Decimal(
                        str(quantity or 0)
                    )
                )

                if quantity < 0:
                    raise ValueError(
                        "Quantity cannot be negative"
                    )

                # -----------------------------------------
                # Discount
                # -----------------------------------------

                discount = Decimal(
                    str(discount or 0)
                    .replace("%", "")
                    .strip()
                )

                if discount < 0:
                    discount = Decimal("0")

                if discount > 100:
                    raise ValueError(
                        "Discount cannot exceed 100%"
                    )

                # -----------------------------------------
                # Amount
                # -----------------------------------------

                amount = Decimal(
                    str(amount or 0)
                )

                if amount < 0:
                    raise ValueError(
                        "Amount cannot be negative"
                    )

                # -----------------------------------------
                # Selling price
                # -----------------------------------------

                if quantity > 0:

                    selling_price = (
                        amount /
                        Decimal(quantity)
                    )

                else:

                    selling_price = amount

                selling_price = selling_price.quantize(
                    Decimal("0.01")
                )

                # -----------------------------------------
                # Calculate MRP
                # -----------------------------------------

                if discount > 0:

                    discount_factor = (
                        Decimal("1")
                        -
                        (
                            discount /
                            Decimal("100")
                        )
                    )

                    if discount_factor <= 0:
                        raise ValueError(
                            "Invalid discount value"
                        )

                    mrp = (
                        selling_price /
                        discount_factor
                    )

                else:

                    mrp = selling_price

                mrp = mrp.quantize(
                    Decimal("0.01")
                )

                # -----------------------------------------
                # Brand
                # -----------------------------------------

                brand_slug = slugify(
                    brand_name
                )

                if brand_slug in brand_cache:

                    brand = brand_cache[
                        brand_slug
                    ]

                else:

                    brand, _ = Brand.objects.get_or_create(
                        slug=brand_slug,
                        defaults={
                            "name": brand_name
                        }
                    )

                    brand_cache[
                        brand_slug
                    ] = brand

                # -----------------------------------------
                # Category
                # -----------------------------------------

                category_slug = slugify(
                    category_name
                )

                if category_slug in category_cache:

                    category = category_cache[
                        category_slug
                    ]

                else:

                    category, _ = Category.objects.get_or_create(
                        slug=category_slug,
                        defaults={
                            "name": category_name
                        }
                    )

                    category_cache[
                        category_slug
                    ] = category

                # -----------------------------------------
                # Create / Update Part
                # -----------------------------------------

                with transaction.atomic():

                    part, created = (
                        Part.objects.update_or_create(

                            part_number=part_number,

                            defaults={
                                "name": name,
                                "brand": brand,
                                "category": category,
                                "description": "",
                                "compatible_with": "",
                                "quantity": quantity,
                                "mrp": mrp,
                                "discount": discount,
                                "is_featured": False,
                                "is_active": True,
                            }
                        )
                    )

                    # Keep this because your model's save()
                    # apparently calculates the price.
                    part.save()

                # -----------------------------------------
                # Counters
                # -----------------------------------------

                if created:

                    created_count += 1

                else:

                    updated_count += 1

            except (
                ValueError,
                InvalidOperation,
                TypeError
            ) as e:

                skipped_count += 1

                # Only retain first 50 errors
                # to prevent memory growth.
                if len(errors) < 50:

                    errors.append(
                        f"Row {row_number}: {str(e)}"
                    )

            except Exception as e:

                skipped_count += 1

                if len(errors) < 50:

                    errors.append(
                        f"Row {row_number}: {str(e)}"
                    )

        # -----------------------------------------
        # Close workbook
        # -----------------------------------------

        workbook.close()
        workbook = None

    except Exception as e:

        if workbook is not None:
            try:
                workbook.close()
            except Exception:
                pass

        messages.error(
            request,
            f"Import failed: {str(e)}"
        )

        return redirect(
            "dashboard:part_list"
        )

    # -----------------------------------------
    # Success message
    # -----------------------------------------

    messages.success(
        request,
        (
            f"Import completed successfully. "
            f"Created: {created_count}, "
            f"Updated: {updated_count}, "
            f"Skipped: {skipped_count}"
        )
    )

    if errors:

        messages.warning(
            request,
            " | ".join(errors)
        )

    return redirect(
        "dashboard:part_list"
    )
from django.core.paginator import Paginator
@admin_required
def part_list_view(request):

    query = request.GET.get(
        "q",
        ""
    ).strip()

    brand_filter = request.GET.get(
        "brand",
        ""
    ).strip()

    category_filter = request.GET.get(
        "category",
        ""
    ).strip()


    parts = (
        Part.objects
        .select_related(
            "brand",
            "category"
        )
        .all()
    )


    # =========================
    # SEARCH
    # =========================

    if query:

        parts = parts.filter(

            Q(name__icontains=query)

            | Q(
                part_number__icontains=query
            )

            | Q(
                compatible_with__icontains=query
            )

            | Q(
                description__icontains=query
            )

        )


    # =========================
    # BRAND FILTER
    # =========================

    if brand_filter:

        parts = parts.filter(
            brand_id=brand_filter
        )


    # =========================
    # CATEGORY FILTER
    # =========================

    if category_filter:

        parts = parts.filter(
            category_id=category_filter
        )


    # =========================
    # PAGINATION
    # =========================

    paginator = Paginator(
        parts,
        10
    )

    page_number = request.GET.get(
        "page"
    )

    page_obj = paginator.get_page(
        page_number
    )


    context = {

        "parts": page_obj,

        "page_obj": page_obj,

        "paginator": paginator,

        "brands": Brand.objects.all(),

        "categories": Category.objects.all(),

        "search_query": query,

        "selected_brand": brand_filter,

        "selected_category": category_filter,
    }


    return render(
        request,
        "dashboard/part_list.html",
        context
    )


@admin_required
def part_create_view(request):

    if request.method == "POST":

        form = PartForm(
            request.POST
        )

        if form.is_valid():

            form.save()

            messages.success(
                request,
                "Spare part added successfully."
            )

            return redirect(
                "dashboard:part_list"
            )

    else:

        form = PartForm()

    return render(
        request,
        "dashboard/part_form.html",
        {
            "form": form,
            "action": "Add",
            "page_title": "Add Spare Part",
        }
    )


@admin_required
def part_update_view(request, pk):

    part = get_object_or_404(
        Part,
        pk=pk
    )

    if request.method == "POST":

        form = PartForm(
            request.POST,
            instance=part
        )

        if form.is_valid():

            form.save()

            messages.success(
                request,
                "Spare part updated successfully."
            )

            return redirect(
                "dashboard:part_list"
            )

    else:

        form = PartForm(
            instance=part
        )

    return render(
        request,
        "dashboard/part_form.html",
        {
            "form": form,
            "action": "Edit",
            "page_title": "Edit Spare Part",
            "part": part,
        }
    )


@admin_required
def part_delete_view(request, pk):

    part = get_object_or_404(
        Part,
        pk=pk
    )

    if request.method == "POST":

        part.delete()

        messages.success(
            request,
            "Spare part deleted successfully."
        )

        return redirect(
            "dashboard:part_list"
        )

    return render(
        request,
        "dashboard/part_confirm_delete.html",
        {
            "part": part
        }
    )


# ============================================================
# BRAND CRUD
# ============================================================

@admin_required
def brand_list_view(request):

    brands = Brand.objects.all()

    return render(
        request,
        "dashboard/brand_list.html",
        {
            "brands": brands
        }
    )


@admin_required
def brand_create_view(request):

    if request.method == "POST":

        form = BrandForm(
            request.POST
        )

        if form.is_valid():

            form.save()

            messages.success(
                request,
                "Brand added successfully."
            )

            return redirect(
                "dashboard:brand_list"
            )

    else:

        form = BrandForm()

    return render(
        request,
        "dashboard/brand_form.html",
        {
            "form": form,
            "action": "Add",
            "page_title": "Add Brand",
        }
    )


@admin_required
def brand_update_view(request, pk):

    brand = get_object_or_404(
        Brand,
        pk=pk
    )

    if request.method == "POST":

        form = BrandForm(
            request.POST,
            instance=brand
        )

        if form.is_valid():

            form.save()

            messages.success(
                request,
                "Brand updated successfully."
            )

            return redirect(
                "dashboard:brand_list"
            )

    else:

        form = BrandForm(
            instance=brand
        )

    return render(
        request,
        "dashboard/brand_form.html",
        {
            "form": form,
            "action": "Edit",
            "page_title": "Edit Brand",
            "brand": brand,
        }
    )


@admin_required
def brand_delete_view(request, pk):

    brand = get_object_or_404(
        Brand,
        pk=pk
    )

    if request.method == "POST":

        brand.delete()

        messages.success(
            request,
            "Brand deleted successfully."
        )

        return redirect(
            "dashboard:brand_list"
        )

    return render(
        request,
        "dashboard/brand_confirm_delete.html",
        {
            "brand": brand
        }
    )


# ============================================================
# CATEGORY CRUD
# ============================================================

@admin_required
def category_list_view(request):

    categories = Category.objects.all()

    return render(
        request,
        "dashboard/category_list.html",
        {
            "categories": categories
        }
    )


@admin_required
def category_create_view(request):

    if request.method == "POST":

        form = CategoryForm(
            request.POST
        )

        if form.is_valid():

            form.save()

            messages.success(
                request,
                "Category added successfully."
            )

            return redirect(
                "dashboard:category_list"
            )

    else:

        form = CategoryForm()

    return render(
        request,
        "dashboard/category_form.html",
        {
            "form": form,
            "action": "Add",
            "page_title": "Add Category",
        }
    )


@admin_required
def category_update_view(request, pk):

    category = get_object_or_404(
        Category,
        pk=pk
    )

    if request.method == "POST":

        form = CategoryForm(
            request.POST,
            instance=category
        )

        if form.is_valid():

            form.save()

            messages.success(
                request,
                "Category updated successfully."
            )

            return redirect(
                "dashboard:category_list"
            )

    else:

        form = CategoryForm(
            instance=category
        )

    return render(
        request,
        "dashboard/category_form.html",
        {
            "form": form,
            "action": "Edit",
            "page_title": "Edit Category",
            "category": category,
        }
    )


@admin_required
def category_delete_view(request, pk):

    category = get_object_or_404(
        Category,
        pk=pk
    )

    if request.method == "POST":

        category.delete()

        messages.success(
            request,
            "Category deleted successfully."
        )

        return redirect(
            "dashboard:category_list"
        )

    return render(
        request,
        "dashboard/category_confirm_delete.html",
        {
            "category": category
        }
    )


# ============================================================
# ENQUIRY CRUD
# ============================================================

@admin_required
def enquiry_list_view(request):

    enquiries = Enquiry.objects.all()

    status = request.GET.get(
        "status",
        ""
    ).strip()

    query = request.GET.get(
        "q",
        ""
    ).strip()

    if status == "pending":

        enquiries = enquiries.filter(
            is_handled=False
        )

    elif status == "handled":

        enquiries = enquiries.filter(
            is_handled=True
        )

    if query:

        enquiries = enquiries.filter(
            Q(name__icontains=query)
            | Q(phone__icontains=query)
            | Q(part__icontains=query)
            | Q(brand__icontains=query)
            | Q(model__icontains=query)
        )

    return render(
        request,
        "dashboard/enquiry_list.html",
        {
            "enquiries": enquiries,
            "search_query": query,
            "selected_status": status,
        }
    )


@admin_required
def enquiry_create_view(request):

    if request.method == "POST":

        form = EnquiryForm(
            request.POST
        )

        if form.is_valid():

            form.save()

            messages.success(
                request,
                "Enquiry created successfully."
            )

            return redirect(
                "dashboard:enquiry_list"
            )

    else:

        form = EnquiryForm()

    return render(
        request,
        "dashboard/enquiry_form.html",
        {
            "form": form,
            "action": "Add",
        }
    )


@admin_required
def enquiry_update_view(request, pk):

    enquiry = get_object_or_404(
        Enquiry,
        pk=pk
    )

    if request.method == "POST":

        form = EnquiryForm(
            request.POST,
            instance=enquiry
        )

        if form.is_valid():

            form.save()

            messages.success(
                request,
                "Enquiry updated successfully."
            )

            return redirect(
                "dashboard:enquiry_list"
            )

    else:

        form = EnquiryForm(
            instance=enquiry
        )

    return render(
        request,
        "dashboard/enquiry_form.html",
        {
            "form": form,
            "action": "Edit",
            "enquiry": enquiry,
        }
    )


@admin_required
def enquiry_delete_view(request, pk):

    enquiry = get_object_or_404(
        Enquiry,
        pk=pk
    )

    if request.method == "POST":

        enquiry.delete()

        messages.success(
            request,
            "Enquiry deleted successfully."
        )

        return redirect(
            "dashboard:enquiry_list"
        )

    return render(
        request,
        "dashboard/enquiry_confirm_delete.html",
        {
            "enquiry": enquiry
        }
    )


@admin_required
def toggle_enquiry_status(request, pk):

    enquiry = get_object_or_404(
        Enquiry,
        pk=pk
    )

    enquiry.is_handled = not enquiry.is_handled

    enquiry.save(
        update_fields=[
            "is_handled"
        ]
    )

    messages.success(
        request,
        "Enquiry status updated."
    )

    return redirect(
        "dashboard:enquiry_list"
    )